# Файл склада для импорта в Google Таблицы: stock/sklad-nosinoski.xlsx
# Запуск из корня проекта: python stock/make-sheet.py
# Колонки A–E не переставлять: их читает Apps Script (A sku | B товар | C размер | D цвет | E остаток)
import json
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

products = json.load(open('shop/src/data/products.json', encoding='utf-8'))
CAT = {
    'women-socks': 'Женские носки', 'men-socks': 'Мужские носки', 'kids': 'Детские носки',
    'sport': 'Спортивные носки', 'warm': 'Тёплые носки', 'soft': 'Мягкая резинка',
    'socks': 'Носки', 'sets': 'Наборы', 'women': 'Женское бельё', 'men': 'Мужское бельё',
}
ORDER = list(CAT)

rows = []
for p in products:
    if not any(not i.endswith('.svg') for i in p['images']):
        continue  # скрытые демо-товары без фото
    for v in p['variants']:
        color = next((c['name'] for c in p['colors'] if c['id'] == v['color']), v['color'])
        rows.append([v['sku'], p['title'], v['size'], color, v['stock'], CAT.get(p['category'], p['category']), p['price']])
rows.sort(key=lambda r: (ORDER.index(next(k for k, n in CAT.items() if n == r[5])), r[1], r[3], r[2]))

wb = Workbook()
ws = wb.active
ws.title = 'stock'
head = ['sku (не менять)', 'товар', 'размер', 'цвет', 'остаток', 'категория', 'цена, ₽']
ws.append(head)
for r in rows:
    ws.append(r)
n = len(rows) + 1

bold = Font(bold=True)
for c in ws[1]:
    c.font = bold
    c.fill = PatternFill('solid', fgColor='DCFF3A')
    c.alignment = Alignment(vertical='center')
grey = PatternFill('solid', fgColor='EEEEEE')
for (c,) in ws.iter_rows(min_row=2, max_row=n, min_col=1, max_col=1):
    c.fill = grey
    c.font = Font(color='888888', size=9)
for (c,) in ws.iter_rows(min_row=2, max_row=n, min_col=5, max_col=5):
    c.font = Font(bold=True, size=12)
    c.alignment = Alignment(horizontal='center')
for col, w in zip('ABCDEFG', (34, 42, 10, 16, 10, 18, 9)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'B2'
ws.auto_filter.ref = f'A1:G{n}'

# 0 — красный, 1–2 — жёлтый
rng = f'E2:E{n}'
ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['0'], fill=PatternFill('solid', fgColor='FFB3B3')))
ws.conditional_formatting.add(rng, CellIsRule(operator='between', formula=['1', '2'], fill=PatternFill('solid', fgColor='FFF2A8')))
dv = DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0', showErrorMessage=True,
                    errorTitle='Остаток', error='Только целое число 0 или больше')
ws.add_data_validation(dv)
dv.add(rng)

# Итого
t = wb.create_sheet('Итого')
t.append(['Показатель', 'Значение'])
t.append(['Всего позиций (модель × цвет × размер)', f'=COUNTA(stock!A2:A{n})'])
t.append(['Всего штук на складе', f'=SUM(stock!E2:E{n})'])
t.append(['Закончились (0)', f'=COUNTIF(stock!E2:E{n},0)'])
t.append(['Заканчиваются (1–2)', f'=COUNTIFS(stock!E2:E{n},">=1",stock!E2:E{n},"<=2")'])
t.append(['Товар на сумму, ₽', f'=SUMPRODUCT(stock!E2:E{n},stock!G2:G{n})'])
t.append([])
t.append(['Штук по категориям', ''])
for name in dict.fromkeys(r[5] for r in rows):
    t.append([name, f'=SUMIF(stock!F2:F{n},"{name}",stock!E2:E{n})'])
for c in t[1]:
    c.font = bold
    c.fill = PatternFill('solid', fgColor='DCFF3A')
t['A8'].font = bold
t.column_dimensions['A'].width = 42
t.column_dimensions['B'].width = 14

# Памятка
h = wb.create_sheet('Как пользоваться')
for line in [
    'Склад «Носи носки» — как пользоваться',
    '',
    'Меняйте только колонку «остаток» на листе stock.',
    'Пришла партия — прибавьте к числу. Продали на Авито или вживую — уменьшите.',
    'Продажи с сайта (когда подключим) спишутся сами и попадут на лист log.',
    'Красный — закончилось, жёлтый — осталось 1–2 шт.',
    'Чтобы найти модель: значок фильтра в заголовке «товар» → выбрать модель.',
    'Колонку sku не меняйте: по ней сайт узнаёт позицию.',
    'Лист «Итого» считает сам: сколько всего пар, что закончилось, сумма товара.',
]:
    h.append([line])
h['A1'].font = Font(bold=True, size=14)
h.column_dimensions['A'].width = 90

out = 'stock/sklad-nosinoski.xlsx'
wb.save(out)
print(f'{out}: {len(rows)} позиций, {len({r[1] for r in rows})} моделей, {sum(r[4] for r in rows)} шт')
